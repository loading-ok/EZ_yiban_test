#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
题库匹配助手 (Question Bank Assistant)
=========================================
功能：
  1. 导入 doc / docx / pdf / xls / xlsx / txt 等格式的题库
  2. 导入 doc / docx / pdf / xls / xlsx / txt 等格式的手册文本
  3. 实时抓取网页内的题目（Selenium 附加到已开启调试端口的 Chrome / Edge）
     ★ 抓取功能独立：不导入题库或手册也能抓取、记录、导出
  4. 依据题库给出正确答案，并从手册中搜索相关答案句子
  5. 手动粘贴题目自动匹配（粘贴框始终可用）
  6. 手动控制开始 / 结束 / 记录当前页
  7. 把抓取到的内容导出为 txt 文件

依赖安装：
  pip install python-docx pdfplumber openpyxl xlrd pandas selenium
  # Windows 下若需读取 .doc，额外安装： pip install pywin32

【免责声明】本脚本仅供个人学习、自测与题库检索研究使用，
请勿用于任何违反考试纪律或平台规定的场景。
"""

import os
import re
import sys
import time
import queue
import threading
import traceback
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import difflib

APP_TITLE = "题库匹配助手"
DISCLAIMER = "仅供学习与自测使用，请勿用于任何违规考试场景。"


# ============================================================
# 一、文本归一化与解析工具
# ============================================================

_NUM_PREFIX = re.compile(
    r'^\s*(?:'
    r'第\s*\d{1,4}\s*[题問问]\s*[\.、:：]?'
    r'|[\(（\[【]\s*\d{1,4}\s*[\)）\]】]'
    r'|\d{1,3}\s*[\.、:：）\)](?!\d)'
    r')\s*'
)


def normalize(s: str) -> str:
    if not s:
        return ""
    s = _NUM_PREFIX.sub('', s.strip())
    s = re.sub(r'[^\w\u4e00-\u9fff]+', '', s)
    return s.lower()


_OPT_SPLIT = re.compile(
    r'(?:^|[\s\u3000])[\(（]?\s*([A-Ha-h])\s*[\)）]?\s*[\.、:：．]\s*'
)


def split_stem_options(text):
    if not text:
        return text, {}
    matches = list(_OPT_SPLIT.finditer(text))
    if not matches:
        return text, {}
    start_i = None
    for i, m in enumerate(matches):
        if m.group(1).upper() == 'A':
            start_i = i
            break
    if start_i is None:
        return text, {}
    ms = matches[start_i:]
    stem = text[:ms[0].start()].strip()
    opts = {}
    for i, m in enumerate(ms):
        end = ms[i + 1].start() if i + 1 < len(ms) else len(text)
        key = m.group(1).upper()
        val = text[m.end():end].strip()
        if key not in opts and val:
            opts[key] = val
    return stem, opts


_RE_ANSWER = re.compile(r'^\s*(?:正确|参考|标准)?\s*答\s*案\s*[:：]?\s*(.*)$')
_RE_OPT1 = re.compile(r'^\s*[\(（]\s*([A-Ha-h])\s*[\)）]\s*[:：\.、]?\s*(.*)$')
_RE_OPT2 = re.compile(r'^\s*([A-Ha-h])\s*[\.、:：．]\s*(.*)$')
_RE_NEWQ = re.compile(r'^\s*(?:第\s*\d{1,4}\s*[题問问]|\d{1,3}\s*[\.、）\)]|[\(（]\s*\d{1,3}\s*[\)）])')


# ============================================================
# 二、数据模型
# ============================================================

class Question:
    __slots__ = ('stem', 'options', 'answer', 'source', '_norm')

    def __init__(self, stem, options=None, answer='', source=''):
        self.stem = (stem or '').strip()
        self.options = dict(options or {})
        self.answer = (answer or '').strip()
        self.source = source
        self._norm = normalize(self.stem)

    @property
    def norm(self):
        return self._norm

    def answer_text(self):
        a = self.answer.strip()
        if not a:
            return '（题库中未提供答案）'
        letters = re.findall(r'[A-Ha-h]', a)
        if letters and all(c.upper() in self.options for c in letters):
            uniq = list(dict.fromkeys(c.upper() for c in letters))
            return '   '.join(f'{c}. {self.options.get(c, "")}'.strip() for c in uniq)
        return a

    def __repr__(self):
        return f'<Question {self.stem[:24]!r} ans={self.answer!r}>'


# ============================================================
# 三、题库文本解析
# ============================================================

def _parse_block(lines, source=''):
    qs = []
    stem_parts = []
    options = {}
    answer = None

    def flush():
        nonlocal stem_parts, options, answer
        stem = re.sub(r'\s+', ' ', ' '.join(p for p in stem_parts if p)).strip()
        if stem:
            s2, o2 = split_stem_options(stem)
            if o2 and not options:
                stem, options = s2, o2
            if stem:
                qs.append(Question(stem, options, answer or '', source))
        stem_parts = []
        options = {}
        answer = None

    for raw in lines:
        line = raw.strip()
        if not line:
            continue

        m = _RE_ANSWER.match(line)
        if m:
            val = m.group(1).strip()
            if val:
                answer = val
                flush()
            continue

        if stem_parts or options:
            mo = _RE_OPT1.match(line) or _RE_OPT2.match(line)
            if mo:
                key = mo.group(1).upper()
                if key not in options:
                    options[key] = mo.group(2).strip()
                    continue

        if (stem_parts or options) and _RE_NEWQ.match(line):
            flush()

        stem_parts.append(line)

    flush()
    return qs


def _parse_tabbed(lines, source=''):
    qs = []
    for line in lines:
        line = line.strip()
        if not line:
            continue
        parts = [p.strip() for p in re.split(r'\t|\|', line)]
        if len(parts) < 2:
            continue

        stem = parts[0]
        if not stem:
            continue

        last = parts[-1]
        if _RE_OPT1.match(last) or _RE_OPT2.match(last):
            ans = ''
            mid = parts[1:]
        else:
            ans = last
            mid = parts[1:-1]

        m = _RE_ANSWER.match(ans)
        if m and m.group(1).strip():
            ans = m.group(1).strip()

        opts = {}
        if mid and any(mid):
            for i, v in enumerate(mid):
                if v:
                    opts[chr(65 + i)] = v

        s2, o2 = split_stem_options(stem)
        if o2 and not opts:
            stem, opts = s2, o2

        if stem:
            qs.append(Question(stem, opts, ans, source))
    return qs


def parse_text_block(text, source=''):
    if not text:
        return []
    text = text.replace('\r\n', '\n').replace('\r', '\n')
    lines = text.split('\n')
    non_empty = [l for l in lines if l.strip()]
    if not non_empty:
        return []

    tabbed = sum(1 for l in non_empty if ('\t' in l or '|' in l))
    if tabbed >= 2 and tabbed >= len(non_empty) * 0.6:
        return _parse_tabbed(non_empty, source)
    return _parse_block(lines, source)


# ============================================================
# 四、文件加载器
# ============================================================

def load_text_file(path):
    for enc in ('utf-8-sig', 'utf-8', 'gbk', 'gb18030', 'big5'):
        try:
            with open(path, 'r', encoding=enc) as f:
                return f.read()
        except (UnicodeDecodeError, LookupError):
            continue
    with open(path, 'r', encoding='utf-8', errors='ignore') as f:
        return f.read()


def load_docx(path):
    try:
        import docx
    except ImportError:
        raise RuntimeError('缺少 python-docx，请执行：pip install python-docx')

    doc = docx.Document(path)
    qs = []

    paras = [p.text for p in doc.paragraphs]
    if any(p.strip() for p in paras):
        qs.extend(parse_text_block('\n'.join(paras), path))

    rows = []
    for tbl in doc.tables:
        for row in tbl.rows:
            rows.append('\t'.join(c.text.strip() for c in row.cells))
    if rows:
        qs.extend(_parse_tabbed([r for r in rows if r.strip()], path))

    seen = set()
    uniq = []
    for q in qs:
        if q.norm and q.norm not in seen:
            seen.add(q.norm)
            uniq.append(q)
    return uniq


def load_doc(path):
    last_err = None

    if sys.platform.startswith('win'):
        try:
            import win32com.client
            import pythoncom
            pythoncom.CoInitialize()
            word = win32com.client.DispatchEx('Word.Application')
            word.Visible = False
            word.DisplayAlerts = 0
            try:
                doc = word.Documents.Open(os.path.abspath(path), ReadOnly=True)
                text = doc.Content.Text
                doc.Close(False)
            finally:
                word.Quit()
            return parse_text_block(text, path)
        except Exception as e:
            last_err = e

    try:
        import subprocess, tempfile, shutil
        soffice = shutil.which('soffice') or shutil.which('libreoffice')
        if soffice:
            outdir = tempfile.mkdtemp(prefix='qbank_')
            subprocess.run(
                [soffice, '--headless', '--convert-to', 'docx', '--outdir', outdir,
                 os.path.abspath(path)],
                check=True, timeout=180,
                stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL
            )
            base = os.path.splitext(os.path.basename(path))[0] + '.docx'
            newp = os.path.join(outdir, base)
            if os.path.exists(newp):
                return load_docx(newp)
    except Exception as e:
        last_err = e

    raise RuntimeError(
        f'无法解析 .doc 文件：{last_err}\n'
        '请安装 Microsoft Word 或 LibreOffice，或将文件另存为 .docx 后重试。'
    )


def _extract_pdf_text(path):
    try:
        import pdfplumber
        with pdfplumber.open(path) as pdf:
            pages = []
            for pg in pdf.pages:
                pages.append(pg.extract_text() or '')
            return '\n'.join(pages)
    except ImportError:
        try:
            from PyPDF2 import PdfReader
            reader = PdfReader(path)
            return '\n'.join((p.extract_text() or '') for p in reader.pages)
        except ImportError:
            raise RuntimeError('缺少 pdfplumber / PyPDF2，请执行：pip install pdfplumber')


def load_pdf(path):
    return parse_text_block(_extract_pdf_text(path), path)


def _is_answer_cell(s):
    s = str(s).strip()
    if not s or len(s) > 30:
        return False
    if re.fullmatch(r'[A-Ha-h]{1,8}', s):
        return True
    if re.fullmatch(r'[A-Ha-h](?:\s*[,，、/]\s*[A-Ha-h])+', s):
        return True
    if re.match(r'^\s*(?:正确|参考|标准)?\s*答\s*案\s*[:：]', s):
        return True
    if s in ('对', '错', '正确', '错误', '√', '×', 'T', 'F', '是', '否'):
        return True
    return False


def _is_header_row(row):
    keywords = ('题', '题干', '问题', '试题', '答案', '选项', '正确答案',
                'question', 'answer', 'option', 'key')
    n = 0
    for c in row:
        c = str(c).strip().lower()
        if not c:
            continue
        if any(k.lower() in c for k in keywords):
            n += 1
    return n >= 2


def _parse_excel_row(cells):
    cells = [str(c).strip() for c in cells]
    while cells and cells[-1] == '':
        cells.pop()
    if not cells:
        return None

    ans = ''
    ans_idx = -1
    for i in range(len(cells) - 1, -1, -1):
        if _is_answer_cell(cells[i]):
            m = _RE_ANSWER.match(cells[i])
            if m and m.group(1).strip():
                ans = m.group(1).strip()
            else:
                ans = cells[i]
            ans_idx = i
            break

    opts = {}
    stem_parts = []
    for i, c in enumerate(cells):
        if i == ans_idx or not c:
            continue
        mo = _RE_OPT1.match(c) or _RE_OPT2.match(c)
        if mo and mo.group(2).strip():
            key = mo.group(1).upper()
            if key not in opts:
                opts[key] = mo.group(2).strip()
                continue
        stem_parts.append(c)

    stem = ' '.join(stem_parts).strip()
    if not stem:
        return None

    s2, o2 = split_stem_options(stem)
    if o2 and not opts:
        stem, opts = s2, o2
    elif o2 and opts:
        for k, v in o2.items():
            opts.setdefault(k, v)
        stem = s2

    return Question(stem, opts, ans)


def load_excel(path):
    try:
        import pandas as pd
    except ImportError:
        raise RuntimeError('缺少 pandas，请执行：pip install pandas openpyxl xlrd')

    ext = os.path.splitext(path)[1].lower()
    engine = 'openpyxl' if ext in ('.xlsx', '.xlsm') else 'xlrd'

    df = None
    try:
        df = pd.read_excel(path, header=None, dtype=str, engine=engine)
    except Exception:
        try:
            df = pd.read_excel(path, header=None, dtype=str)
        except Exception as e:
            raise RuntimeError(f'读取 Excel 失败：{e}')

    raw_rows = df.fillna('').values.tolist()

    max_cols = max((sum(1 for c in r if str(c).strip()) for r in raw_rows), default=0)
    if max_cols <= 1:
        text = '\n'.join(
            ' '.join(str(c).strip() for c in r if str(c).strip())
            for r in raw_rows
        )
        return parse_text_block(text, path)

    data_rows = raw_rows
    if raw_rows and _is_header_row(raw_rows[0]):
        data_rows = raw_rows[1:]

    qs = []
    for r in data_rows:
        q = _parse_excel_row(r)
        if q:
            qs.append(q)

    seen = set()
    uniq = []
    for q in qs:
        if q.norm and q.norm not in seen:
            seen.add(q.norm)
            uniq.append(q)
    return uniq


def load_bank(path):
    ext = os.path.splitext(path)[1].lower()
    if ext in ('.txt', '.text', '.md', '.csv'):
        return parse_text_block(load_text_file(path), path)
    if ext == '.docx':
        return load_docx(path)
    if ext == '.doc':
        return load_doc(path)
    if ext == '.pdf':
        return load_pdf(path)
    if ext in ('.xlsx', '.xls', '.xlsm'):
        return load_excel(path)
    raise ValueError(f'不支持的格式：{ext}')


# ============================================================
# 五、手册文本提取与搜索
# ============================================================

def extract_text_from_file(path):
    ext = os.path.splitext(path)[1].lower()
    if ext in ('.txt', '.text', '.md', '.csv'):
        return load_text_file(path)
    if ext == '.docx':
        try:
            import docx
        except ImportError:
            raise RuntimeError('缺少 python-docx，请执行：pip install python-docx')
        doc = docx.Document(path)
        parts = []
        for p in doc.paragraphs:
            if p.text.strip():
                parts.append(p.text)
        for tbl in doc.tables:
            for row in tbl.rows:
                parts.append(' '.join(c.text.strip() for c in row.cells))
        return '\n'.join(parts)
    if ext == '.doc':
        if sys.platform.startswith('win'):
            try:
                import win32com.client
                import pythoncom
                pythoncom.CoInitialize()
                word = win32com.client.DispatchEx('Word.Application')
                word.Visible = False
                word.DisplayAlerts = 0
                try:
                    doc = word.Documents.Open(os.path.abspath(path), ReadOnly=True)
                    text = doc.Content.Text
                    doc.Close(False)
                finally:
                    word.Quit()
                return text
            except Exception:
                pass
        try:
            import subprocess, tempfile, shutil
            soffice = shutil.which('soffice') or shutil.which('libreoffice')
            if soffice:
                outdir = tempfile.mkdtemp(prefix='manual_')
                subprocess.run(
                    [soffice, '--headless', '--convert-to', 'docx', '--outdir', outdir,
                     os.path.abspath(path)],
                    check=True, timeout=180,
                    stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL
                )
                base = os.path.splitext(os.path.basename(path))[0] + '.docx'
                newp = os.path.join(outdir, base)
                if os.path.exists(newp):
                    return extract_text_from_file(newp)
        except Exception:
            pass
        raise RuntimeError('无法解析 .doc 手册，请安装 Word 或 LibreOffice')
    if ext == '.pdf':
        return _extract_pdf_text(path)
    if ext in ('.xlsx', '.xls', '.xlsm'):
        try:
            import pandas as pd
        except ImportError:
            raise RuntimeError('缺少 pandas，请执行：pip install pandas openpyxl xlrd')
        df = pd.read_excel(path, header=None, dtype=str)
        rows = []
        for r in df.fillna('').values.tolist():
            rows.append(' '.join(str(c).strip() for c in r if str(c).strip()))
        return '\n'.join(rows)
    return load_text_file(path)


class ManualSearcher:
    """在手册文本中搜索与题目最相关的句子"""

    def __init__(self, text):
        self.sentences = self._split_sentences(text)
        self.norm_sentences = [(normalize(s), s) for s in self.sentences if len(s) > 5]

    @staticmethod
    def _split_sentences(text):
        text = re.sub(r'[\r\n]+', '\n', text)
        parts = re.split(r'(?<=[。！？；\.!?;])\s*|\n+', text)
        return [p.strip() for p in parts if p.strip()]

    def search(self, query, top_n=3):
        if not self.norm_sentences or not query:
            return []
        q_norm = normalize(query)
        if len(q_norm) < 4:
            return []

        scored = []
        for norm_s, orig_s in self.norm_sentences:
            if not set(q_norm) & set(norm_s):
                continue
            ratio = difflib.SequenceMatcher(None, q_norm, norm_s).ratio()
            if q_norm in norm_s:
                ratio = max(ratio, 0.9)
            scored.append((ratio, orig_s))
        scored.sort(key=lambda x: -x[0])

        result = []
        seen = set()
        for _, s in scored:
            if _ <= 0.3:
                break
            if s in seen:
                continue
            seen.add(s)
            result.append(s)
            if len(result) >= top_n:
                break
        return result


# ============================================================
# 六、题库匹配器
# ============================================================

class Matcher:
    """把题库建成索引，用「页面文本包含题目特征串」的方式匹配。"""

    MIN_STEM_LEN = 6

    def __init__(self, questions):
        self.questions = list(questions)
        self.by_norm = {}
        for q in self.questions:
            n = q.norm
            if len(n) >= self.MIN_STEM_LEN and n not in self.by_norm:
                self.by_norm[n] = q

        self.probes = []
        for n, q in self.by_norm.items():
            probe = n[:min(len(n), 40)]
            self.probes.append((probe, n, q))
        self.probes.sort(key=lambda x: -len(x[0]))

    def match_all(self, page_text):
        np = normalize(page_text)
        if not np:
            return []
        results = []
        seen = set()
        for probe, n, q in self.probes:
            if q.norm in seen:
                continue
            if probe in np:
                score = 1.0 if n in np else 0.85
                results.append((q, score))
                seen.add(q.norm)
        return results


# ============================================================
# 七、网页抓取
# ============================================================

class WebGrabber:
    def __init__(self, port=9222, browser='chrome'):
        self.port = port
        self.browser = browser.lower()
        self.driver = None

    def attach(self):
        try:
            from selenium import webdriver
        except ImportError:
            raise RuntimeError('缺少 selenium，请执行：pip install selenium')

        if self.browser == 'edge':
            opts = webdriver.EdgeOptions()
            opts.add_experimental_option('debuggerAddress', f'127.0.0.1:{self.port}')
            self.driver = webdriver.Edge(options=opts)
        else:
            opts = webdriver.ChromeOptions()
            opts.add_experimental_option('debuggerAddress', f'127.0.0.1:{self.port}')
            self.driver = webdriver.Chrome(options=opts)
        return True

    def page_text(self):
        if not self.driver:
            return ''
        try:
            return self.driver.execute_script(
                'return document.body ? document.body.innerText : "";'
            ) or ''
        except Exception:
            return ''

    def close(self):
        self.driver = None


# ============================================================
# 八、答案展示窗
# ============================================================

class AnswerWindow:
    def __init__(self, master):
        self.win = tk.Toplevel(master)
        self.win.title('匹配结果')
        self.win.attributes('-topmost', True)
        self.win.geometry('540x460+60+60')
        self.win.protocol('WM_DELETE_WINDOW', self.hide)

        frm = ttk.Frame(self.win, padding=8)
        frm.pack(fill='both', expand=True)

        self.txt = tk.Text(frm, wrap='word', font=('Microsoft YaHei', 11),
                           bg='#fffdf5', relief='flat')
        sb = ttk.Scrollbar(frm, command=self.txt.yview)
        self.txt.configure(yscrollcommand=sb.set)
        sb.pack(side='right', fill='y')
        self.txt.pack(side='left', fill='both', expand=True)

        self.txt.tag_configure('title',
                               foreground='#333',
                               font=('Microsoft YaHei', 10, 'bold'))
        self.txt.tag_configure('ans',
                               foreground='#c0392b',
                               font=('Microsoft YaHei', 16, 'bold'))
        self.txt.tag_configure('dim',
                               foreground='#888',
                               font=('Microsoft YaHei', 9))
        self.txt.tag_configure('manual',
                               foreground='#1a5276',
                               font=('Microsoft YaHei', 10))
        self.txt.tag_configure('count',
                               foreground='#0066aa',
                               font=('Microsoft YaHei', 10, 'bold'))
        self.txt.configure(state='disabled')

    def show(self):
        self.win.deiconify()
        self.win.lift()

    def hide(self):
        self.win.withdraw()

    def set_result(self, all_matches, manual_snippets=None, raw_preview=None):
        """all_matches: [(q, score), ...]  可能为空
           manual_snippets: 手册参考句子
           raw_preview: 未命中时展示的原始页面文本（前几百字）
        """
        self.txt.configure(state='normal')
        self.txt.delete('1.0', 'end')

        if all_matches:
            self.txt.insert('end',
                            f'本次命中 {len(all_matches)} 道题\n\n', 'count')
            for idx, (q, score) in enumerate(all_matches, 1):
                self.txt.insert('end', f'【第 {idx} 题】\n', 'title')
                self.txt.insert('end', q.stem + '\n')
                if q.options:
                    for k in sorted(q.options):
                        self.txt.insert('end', f'  {k}. {q.options[k]}\n')
                self.txt.insert('end', '答案：', 'title')
                self.txt.insert('end', q.answer_text() + '\n', 'ans')
                self.txt.insert('end', f'匹配度：{score:.0%}\n\n', 'dim')
        else:
            self.txt.insert('end', '未在题库中找到匹配的题目\n\n', 'title')
            if raw_preview:
                self.txt.insert('end', '页面内容预览：\n', 'title')
                self.txt.insert('end', raw_preview[:600], 'dim')

        if manual_snippets:
            self.txt.insert('end', '\n── 手册参考 ──\n', 'title')
            for i, snip in enumerate(manual_snippets, 1):
                self.txt.insert('end', f'{i}. {snip}\n', 'manual')

        self.txt.configure(state='disabled')
        self.txt.see('1.0')


# ============================================================
# 九、主界面
# ============================================================

class App:
    def __init__(self, root):
        self.root = root
        root.title(APP_TITLE)
        root.geometry('780x660')
        root.minsize(720, 600)

        self.matcher = None
        self.manual_searcher = None
        self.grabber = WebGrabber()
        self.running = False
        self.stop_evt = threading.Event()
        self.worker = None
        self.msg_queue = queue.Queue()
        self._last_key = None
        self._debounce_id = None

        # 抓取记录
        self.captured_items = []
        self._captured_keys = set()

        self._build_ui()

        self.answer_win = AnswerWindow(root)
        self.answer_win.hide()

        self.root.after(200, self._poll_queue)
        self.root.protocol('WM_DELETE_WINDOW', self.on_close)

    # ---------- 界面构建 ----------
    def _build_ui(self):
        pad = dict(padx=8, pady=4)

        # ---- 题库区 ----
        lf_bank = ttk.LabelFrame(self.root, text=' 题库（可只导入题库，也可完全跳过） ')
        lf_bank.pack(fill='x', **pad)

        self.bank_path_var = tk.StringVar(value='（尚未导入）')
        ttk.Entry(lf_bank, textvariable=self.bank_path_var,
                  state='readonly').pack(side='left', fill='x', expand=True,
                                         padx=(8, 4), pady=8)

        ttk.Button(lf_bank, text='导入题库',
                   command=self.on_import_bank).pack(side='left', padx=4, pady=8)
        ttk.Button(lf_bank, text='清空',
                   command=self.on_clear_bank).pack(side='left', padx=(0, 8), pady=8)

        self.bank_info_var = tk.StringVar(value='题目数：0')
        ttk.Label(lf_bank, textvariable=self.bank_info_var,
                  foreground='#0066aa').pack(side='left', padx=(0, 8))

        # ---- 手册区 ----
        lf_manual = ttk.LabelFrame(self.root, text=' 手册（可选，用于搜索答案） ')
        lf_manual.pack(fill='x', **pad)

        self.manual_path_var = tk.StringVar(value='（尚未导入）')
        ttk.Entry(lf_manual, textvariable=self.manual_path_var,
                  state='readonly').pack(side='left', fill='x', expand=True,
                                         padx=(8, 4), pady=8)

        ttk.Button(lf_manual, text='导入手册',
                   command=self.on_import_manual).pack(side='left', padx=4, pady=8)
        ttk.Button(lf_manual, text='清空',
                   command=self.on_clear_manual).pack(side='left', padx=(0, 8), pady=8)

        self.manual_info_var = tk.StringVar(value='句子数：0')
        ttk.Label(lf_manual, textvariable=self.manual_info_var,
                  foreground='#0066aa').pack(side='left', padx=(0, 8))

        # ---- 数据源区 ----
        lf_src = ttk.LabelFrame(self.root, text=' 数据源 ')
        lf_src.pack(fill='x', **pad)

        self.source_var = tk.StringVar(value='web')

        row1 = ttk.Frame(lf_src)
        row1.pack(fill='x', padx=8, pady=(8, 2))
        ttk.Radiobutton(row1, text='网页实时抓取（Selenium）',
                        variable=self.source_var, value='web',
                        command=self._on_source_change).pack(side='left')
        ttk.Label(row1, text='   浏览器：').pack(side='left')
        self.browser_var = tk.StringVar(value='chrome')
        ttk.Combobox(row1, textvariable=self.browser_var,
                     values=['chrome', 'edge'], width=8,
                     state='readonly').pack(side='left')
        ttk.Label(row1, text=' 端口：').pack(side='left')
        self.port_var = tk.StringVar(value='9222')
        ttk.Entry(row1, textvariable=self.port_var, width=8).pack(side='left')
        self.btn_attach = ttk.Button(row1, text='连接浏览器', command=self.on_attach)
        self.btn_attach.pack(side='left', padx=8)

        row2 = ttk.Frame(lf_src)
        row2.pack(fill='x', padx=8, pady=2)
        ttk.Radiobutton(row2, text='手动粘贴（把题目粘到下方框内即可自动匹配）',
                        variable=self.source_var, value='manual',
                        command=self._on_source_change).pack(side='left')

        self.conn_var = tk.StringVar(value='未连接')
        ttk.Label(lf_src, textvariable=self.conn_var,
                  foreground='#888').pack(anchor='w', padx=12, pady=(0, 8))

        # ---- 粘贴区 ----
        lf_paste = ttk.LabelFrame(self.root, text=' 手动粘贴区（任何模式下都可粘贴） ')
        lf_paste.pack(fill='both', expand=True, **pad)

        self.paste_txt = tk.Text(lf_paste, height=6, wrap='word',
                                 font=('Microsoft YaHei', 10),
                                 undo=True)
        self.paste_txt.pack(fill='both', expand=True, padx=8, pady=8)

        self.paste_menu = tk.Menu(self.root, tearoff=0)
        self.paste_menu.add_command(label='粘贴', command=self._paste_clipboard)
        self.paste_menu.add_command(label='复制', command=self._copy_selection)
        self.paste_menu.add_command(label='剪切', command=self._cut_selection)
        self.paste_menu.add_separator()
        self.paste_menu.add_command(label='全选', command=self._select_all)
        self.paste_menu.add_command(label='清空', command=self._clear_paste)

        self.paste_txt.bind('<Button-3>', self._show_paste_menu)
        self.paste_txt.bind('<Button-2>', self._show_paste_menu)
        self.paste_txt.bind('<Control-v>', lambda e: self._paste_clipboard())
        self.paste_txt.bind('<Control-V>', lambda e: self._paste_clipboard())
        self.paste_txt.bind('<Control-a>', lambda e: self._select_all())
        self.paste_txt.bind('<Control-A>', lambda e: self._select_all())

        self.paste_txt.bind('<KeyRelease>', self._on_paste_change)
        self.paste_txt.bind('<<Paste>>', self._on_paste_change)
        self.paste_txt.bind('<<Cut>>', self._on_paste_change)
        self.paste_txt.bind('<<Clear>>', self._on_paste_change)

        # ---- 控制区 ----
        lf_ctrl = ttk.LabelFrame(self.root, text=' 控制 ')
        lf_ctrl.pack(fill='x', **pad)

        self.btn_start = ttk.Button(lf_ctrl, text='▶  开始', command=self.on_start)
        self.btn_start.pack(side='left', padx=10, pady=10, ipadx=18)

        self.btn_stop = ttk.Button(lf_ctrl, text='■  结束',
                                   command=self.on_stop, state='disabled')
        self.btn_stop.pack(side='left', padx=4, pady=10, ipadx=18)

        self.btn_record = ttk.Button(lf_ctrl, text='记录当前页',
                                     command=self.on_record_current)
        self.btn_record.pack(side='left', padx=4, pady=10)

        self.btn_export = ttk.Button(lf_ctrl, text='导出题目',
                                     command=self.on_export_captured)
        self.btn_export.pack(side='left', padx=4, pady=10)

        ttk.Button(lf_ctrl, text='清空记录',
                   command=self.on_clear_captured).pack(side='left', padx=4, pady=10)

        ttk.Button(lf_ctrl, text='显示/隐藏答案窗',
                   command=self._toggle_answer_win).pack(side='right', padx=6, pady=10)

        self.captured_var = tk.StringVar(value='抓取记录：0 条')
        ttk.Label(lf_ctrl, textvariable=self.captured_var,
                  foreground='#0066aa').pack(side='right', padx=6)

        # ---- 状态栏 ----
        self.status_var = tk.StringVar(value='就绪')
        ttk.Label(self.root, textvariable=self.status_var,
                  relief='sunken', anchor='w').pack(fill='x', side='bottom')

        self._on_source_change()

    # ---------- 剪贴板 / 右键菜单 ----------
    def _show_paste_menu(self, event):
        try:
            self.paste_menu.tk_popup(event.x_root, event.y_root)
        finally:
            try:
                self.paste_menu.grab_release()
            except Exception:
                pass
        return 'break'

    def _paste_clipboard(self):
        try:
            text = self.root.clipboard_get()
        except tk.TclError:
            return 'break'
        try:
            self.paste_txt.delete('sel.first', 'sel.last')
        except tk.TclError:
            pass
        self.paste_txt.insert('insert', text)
        self._on_paste_change()
        return 'break'

    def _copy_selection(self):
        try:
            text = self.paste_txt.get('sel.first', 'sel.last')
        except tk.TclError:
            return 'break'
        self.root.clipboard_clear()
        self.root.clipboard_append(text)
        return 'break'

    def _cut_selection(self):
        self._copy_selection()
        try:
            self.paste_txt.delete('sel.first', 'sel.last')
        except tk.TclError:
            pass
        self._on_paste_change()
        return 'break'

    def _select_all(self):
        self.paste_txt.tag_add('sel', '1.0', 'end')
        self.paste_txt.mark_set('insert', '1.0')
        return 'break'

    def _clear_paste(self):
        self.paste_txt.delete('1.0', 'end')
        return 'break'

    # ---------- 事件 ----------
    def _on_source_change(self):
        web = (self.source_var.get() == 'web')
        self.btn_attach.configure(state='normal' if web else 'disabled')
        try:
            self.paste_txt.configure(bg='#ffffff' if not web else '#f7fbff')
        except Exception:
            pass

    def _on_paste_change(self, event=None):
        if self._debounce_id:
            try:
                self.root.after_cancel(self._debounce_id)
            except Exception:
                pass
        self._debounce_id = self.root.after(350, self._do_manual_match)

    def _do_manual_match(self):
        self._debounce_id = None
        text = self.paste_txt.get('1.0', 'end').strip()
        if not text:
            return
        self.process_text(text, source='manual')

    # ---------- 题库导入 ----------
    def on_import_bank(self):
        path = filedialog.askopenfilename(
            title='选择题库文件',
            filetypes=[
                ('所有支持的格式',
                 '*.txt *.md *.csv *.doc *.docx *.pdf *.xls *.xlsx *.xlsm'),
                ('文本文件', '*.txt *.md *.csv'),
                ('Word 文档', '*.doc *.docx'),
                ('PDF 文件', '*.pdf'),
                ('Excel 表格', '*.xls *.xlsx *.xlsm'),
                ('所有文件', '*.*'),
            ]
        )
        if not path:
            return
        self.status_var.set(f'正在解析题库：{os.path.basename(path)} ...')
        self.root.update_idletasks()
        try:
            qs = load_bank(path)
        except Exception as e:
            messagebox.showerror('导入失败',
                                 f'{e}\n\n{traceback.format_exc(limit=2)}')
            self.status_var.set('题库导入失败')
            return

        if not qs:
            messagebox.showwarning(
                '提示',
                '未能从文件中解析出任何题目。\n\n'
                '请确认题库文本包含题干与答案，常见格式如：\n'
                '  1. 中国的首都是？\n'
                '  A. 北京   B. 上海\n'
                '  答案：A\n'
                '或每行「题干<TAB>答案」。'
            )
            self.status_var.set('未解析到题目')
            return

        self.matcher = Matcher(qs)
        self.bank_path_var.set(path)
        self.bank_info_var.set(f'题目数：{len(self.matcher.by_norm)}')
        self.status_var.set(f'题库导入成功，共 {len(self.matcher.by_norm)} 道题（可匹配）')

    def on_clear_bank(self):
        self.matcher = None
        self.bank_path_var.set('（尚未导入）')
        self.bank_info_var.set('题目数：0')
        self.status_var.set('题库已清空')

    # ---------- 手册导入 ----------
    def on_import_manual(self):
        path = filedialog.askopenfilename(
            title='选择手册文件',
            filetypes=[
                ('所有支持的格式',
                 '*.txt *.md *.doc *.docx *.pdf *.xls *.xlsx *.xlsm'),
                ('文本文件', '*.txt *.md'),
                ('Word 文档', '*.doc *.docx'),
                ('PDF 文件', '*.pdf'),
                ('Excel 表格', '*.xls *.xlsx *.xlsm'),
                ('所有文件', '*.*'),
            ]
        )
        if not path:
            return
        self.status_var.set(f'正在解析手册：{os.path.basename(path)} ...')
        self.root.update_idletasks()
        try:
            text = extract_text_from_file(path)
        except Exception as e:
            messagebox.showerror('导入手册失败',
                                 f'{e}\n\n{traceback.format_exc(limit=2)}')
            self.status_var.set('手册导入失败')
            return

        if not text.strip():
            messagebox.showwarning('提示', '手册内容为空')
            self.status_var.set('手册内容为空')
            return

        self.manual_searcher = ManualSearcher(text)
        self.manual_path_var.set(path)
        self.manual_info_var.set(f'句子数：{len(self.manual_searcher.sentences)}')
        self.status_var.set(
            f'手册导入成功，共 {len(self.manual_searcher.sentences)} 个句子'
        )

    def on_clear_manual(self):
        self.manual_searcher = None
        self.manual_path_var.set('（尚未导入）')
        self.manual_info_var.set('句子数：0')
        self.status_var.set('手册已清空')

    # ---------- 浏览器连接 ----------
    def on_attach(self):
        try:
            port = int(self.port_var.get().strip() or '9222')
        except ValueError:
            messagebox.showwarning('提示', '端口号必须是数字')
            return
        browser = self.browser_var.get()
        self.grabber = WebGrabber(port, browser)
        try:
            self.grabber.attach()
        except Exception as e:
            self.conn_var.set('未连接')
            messagebox.showerror(
                '连接失败',
                f'{e}\n\n请先以调试模式启动浏览器：\n'
                f'Chrome:\n  chrome.exe --remote-debugging-port={port} '
                f'--user-data-dir="C:\\chrome_debug"\n\n'
                f'Edge:\n  msedge.exe --remote-debugging-port={port} '
                f'--user-data-dir="C:\\edge_debug"\n\n'
                '然后用这个浏览器打开题目页面，再点击「连接浏览器」。'
            )
            return
        self.conn_var.set(f'已连接 127.0.0.1:{port} ({browser})')
        self.status_var.set('浏览器连接成功')

    # ---------- 处理文本（核心）----------
    def process_text(self, text, source='web'):
        if not text:
            return

        # 1. 匹配题库
        matches = []
        if self.matcher:
            matches = self.matcher.match_all(text)

        # 2. 记录抓取结果（命中的记结构化题目，未命中的记原始文本）
        self._record_capture(text, matches)

        # 3. 展示
        sorted_matches = sorted(matches, key=lambda x: -x[1])
        top_q = sorted_matches[0][0] if sorted_matches else None

        manual_snips = []
        if self.manual_searcher:
            query = top_q.stem if top_q else text
            manual_snips = self.manual_searcher.search(query)

        raw_preview = None if sorted_matches else text
        self.answer_win.show()
        self.answer_win.set_result(sorted_matches, manual_snips, raw_preview)

        # 4. 状态栏
        parts = []
        if sorted_matches:
            parts.append(f'命中 {len(sorted_matches)} 道题')
        else:
            parts.append('未命中题库')
        if manual_snips:
            parts.append(f'手册参考 {len(manual_snips)} 条')
        parts.append(f'已记录 {len(self.captured_items)} 条')
        self.status_var.set('  |  '.join(parts))

    def _clean_page_text(self, text):
        """清理页面文本：去空行、过短行、首尾空白"""
        if not text:
            return ''
        lines = [l.strip() for l in text.split('\n')]
        lines = [l for l in lines if len(l) >= 2]
        return '\n'.join(lines)

    def _record_capture(self, raw_text, matches):
        """记录本次抓取结果。
        - 若命中题库：记录结构化题目（去重）
        - 若未命中：记录原始页面文本（去重）
        """
        if matches:
            for q, score in matches:
                key = ('Q', q.norm)
                if key in self._captured_keys:
                    continue
                self._captured_keys.add(key)
                self.captured_items.append({
                    'type': 'question',
                    'stem': q.stem,
                    'options': dict(q.options),
                    'answer': q.answer,
                })
        else:
            cleaned = self._clean_page_text(raw_text)
            if not cleaned:
                return
            # 用归一化后的前 300 字符做去重 key
            norm_key = normalize(cleaned)[:300]
            if not norm_key:
                return
            key = ('T', norm_key)
            if key in self._captured_keys:
                return
            self._captured_keys.add(key)
            self.captured_items.append({
                'type': 'text',
                'stem': cleaned,
                'options': {},
                'answer': '',
            })
        self.captured_var.set(f'抓取记录：{len(self.captured_items)} 条')

    # ---------- 手动记录当前页 ----------
    def on_record_current(self):
        if self.source_var.get() == 'web':
            if not self.grabber.driver:
                messagebox.showwarning('提示', '请先点击「连接浏览器」')
                return
            text = self.grabber.page_text()
            if not text.strip():
                messagebox.showinfo('提示', '页面内容为空')
                return
            self.process_text(text, source='web')
            self.status_var.set(
                f'已手动记录，当前共 {len(self.captured_items)} 条'
            )
        else:
            text = self.paste_txt.get('1.0', 'end').strip()
            if not text:
                messagebox.showinfo('提示', '粘贴框为空')
                return
            self.process_text(text, source='manual')
            self.status_var.set(
                f'已手动记录，当前共 {len(self.captured_items)} 条'
            )

    # ---------- 导出 ----------
    def on_export_captured(self):
        if not self.captured_items:
            messagebox.showinfo('提示', '还没有抓取到任何内容，无法导出。')
            return

        default_name = time.strftime('captured_%Y%m%d_%H%M%S.txt')
        path = filedialog.asksaveasfilename(
            title='导出抓取的内容',
            defaultextension='.txt',
            initialfile=default_name,
            filetypes=[('文本文件', '*.txt'), ('所有文件', '*.*')]
        )
        if not path:
            return

        try:
            lines = []
            lines.append(f'# 抓取导出  {time.strftime("%Y-%m-%d %H:%M:%S")}')
            lines.append(f'# 共 {len(self.captured_items)} 条')
            lines.append('')

            for i, item in enumerate(self.captured_items, 1):
                lines.append(f'===== 第 {i} 条 =====')
                lines.append(item['stem'])
                if item.get('options'):
                    for k in sorted(item['options']):
                        lines.append(f'{k}. {item["options"][k]}')
                if item.get('answer'):
                    lines.append(f'答案：{item["answer"]}')
                lines.append('')

            with open(path, 'w', encoding='utf-8') as f:
                f.write('\n'.join(lines))

            messagebox.showinfo(
                '导出成功',
                f'已导出 {len(self.captured_items)} 条到：\n{path}'
            )
            self.status_var.set(
                f'已导出 {len(self.captured_items)} 条到 {os.path.basename(path)}'
            )
        except Exception as e:
            messagebox.showerror('导出失败',
                                 f'{e}\n\n{traceback.format_exc(limit=2)}')

    def on_clear_captured(self):
        if not self.captured_items:
            return
        if messagebox.askyesno(
            '确认',
            f'确定清空已记录的 {len(self.captured_items)} 条抓取内容吗？'
        ):
            self.captured_items.clear()
            self._captured_keys.clear()
            self.captured_var.set('抓取记录：0 条')
            self.status_var.set('已清空抓取记录')

    # ---------- 开始 / 结束 ----------
    def on_start(self):
        if self.running:
            return

        if self.source_var.get() == 'web':
            if not self.grabber.driver:
                messagebox.showwarning('提示', '请先点击「连接浏览器」')
                return
            # ★ 移除对 matcher / manual_searcher 的强制要求
            self.stop_evt.clear()
            self._last_key = None
            self.running = True
            self.worker = threading.Thread(target=self._loop, daemon=True)
            self.worker.start()
            if self.matcher or self.manual_searcher:
                self.status_var.set('运行中（网页抓取 + 匹配）…')
            else:
                self.status_var.set('运行中（纯抓取模式，仅记录页面文本）…')
        else:
            self.status_var.set('运行中（手动粘贴）…')

        self.answer_win.show()
        self.btn_start.configure(state='disabled')
        self.btn_stop.configure(state='normal')

    def on_stop(self):
        self.running = False
        self.stop_evt.set()
        self.btn_start.configure(state='normal')
        self.btn_stop.configure(state='disabled')
        self.status_var.set('已停止')

    def _toggle_answer_win(self):
        if self.answer_win.win.state() == 'withdrawn':
            self.answer_win.show()
        else:
            self.answer_win.hide()

    # ---------- 后台轮询 ----------
    def _loop(self):
        while not self.stop_evt.is_set():
            try:
                text = self.grabber.page_text()
                if text:
                    key = hash(text)
                    if key != self._last_key:
                        self._last_key = key
                        self.msg_queue.put(('text', text))
            except Exception as e:
                self.msg_queue.put(('error', str(e)))
            self.stop_evt.wait(0.8)

    def _poll_queue(self):
        try:
            while True:
                msg = self.msg_queue.get_nowait()
                if msg[0] == 'text':
                    _, text = msg
                    self.process_text(text, source='web')
                elif msg[0] == 'error':
                    self.status_var.set(f'抓取出错：{msg[1]}')
        except queue.Empty:
            pass
        self.root.after(200, self._poll_queue)

    def on_close(self):
        self.running = False
        self.stop_evt.set()
        try:
            self.grabber.close()
        except Exception:
            pass
        self.root.destroy()


# ============================================================
# 十、入口
# ============================================================

def main():
    print(__doc__)
    print('=' * 60)
    print(DISCLAIMER)
    print('=' * 60)

    root = tk.Tk()
    try:
        root.call('tk', 'scaling', 1.2)
    except Exception:
        pass
    App(root)
    root.mainloop()


if __name__ == '__main__':
    main()